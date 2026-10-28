import streamlit as st
import fitz  # PyMuPDF
from pdf2docx import Converter
import tempfile
import os
import random
import string
import traceback
from io import BytesIO
from PIL import Image, ImageEnhance
from streamlit_image_coordinates import streamlit_image_coordinates
import pdfplumber
import pandas as pd
from pptx import Presentation
from pptx.util import Inches

# PyHanko y Cryptography
from pyhanko.sign import signers
from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
from cryptography.hazmat.primitives.serialization import pkcs12

st.set_page_config(page_title="Gestor PDF", page_icon="📄", layout="wide") 

# --- INICIALIZACIÓN DE VARIABLES DE ESTADO ---
if "firma_x" not in st.session_state: st.session_state.firma_x = 50.0
if "firma_y" not in st.session_state: st.session_state.firma_y = 50.0
if "last_click_firma" not in st.session_state: st.session_state.last_click_firma = None

if "texto_x" not in st.session_state: st.session_state.texto_x = 50.0
if "texto_y" not in st.session_state: st.session_state.texto_y = 50.0
if "last_click_texto" not in st.session_state: st.session_state.last_click_texto = None

if "textos_agregados" not in st.session_state: st.session_state.textos_agregados = []
if "pdf_actual_form" not in st.session_state: st.session_state.pdf_actual_form = None

def update_firma_pos():
    st.session_state.firma_x = st.session_state.slider_firma_x
    st.session_state.firma_y = st.session_state.slider_firma_y

def update_texto_pos():
    st.session_state.texto_x = st.session_state.slider_texto_x
    st.session_state.texto_y = st.session_state.slider_texto_y

st.sidebar.title("Módulos PDF")
opcion = st.sidebar.radio("Selecciona una herramienta:", 
                          ["Comprimir PDF", 
                           "Unir y Reorganizar", 
                           "Dividir o Quitar Hojas",
                           "Firmar PDF", 
                           "Editar y Rellenar PDF", 
                           "Marca de Agua",
                           "Proteger / Desbloquear PDF",
                           "Convertir a Word, Excel, PPT"])

# ==========================================
# MÓDULO 1: COMPRIMIR PDF
# ==========================================
if opcion == "Comprimir PDF":
    st.header("🗜️ Comprimir PDF")
    archivo_pdf = st.file_uploader("Sube tu archivo PDF", type=["pdf"])
    if archivo_pdf is not None:
        nivel = st.radio("Nivel de compresión:", ["Medio (Buena calidad)", "Máximo (Calidad media)"])
        if st.button("Comprimir"):
            with st.spinner("Comprimiendo..."):
                doc = fitz.open(stream=archivo_pdf.getvalue(), filetype="pdf")
                garbage_level = 3 if nivel == "Medio (Buena calidad)" else 4
                salida_pdf = BytesIO()
                doc.save(salida_pdf, garbage=garbage_level, deflate=True)
                doc.close()
                st.success("¡Compresión finalizada!")
                st.download_button(label="Descargar PDF Comprimido", data=salida_pdf.getvalue(), file_name=f"comprimido_{archivo_pdf.name}", mime="application/pdf")

# ==========================================
# MÓDULO 2: UNIR Y REORGANIZAR
# ==========================================
elif opcion == "Unir y Reorganizar":
    st.header("🔗 Unir y Reorganizar PDFs")
    archivos_pdf = st.file_uploader("Sube los PDFs a unir", type=["pdf"], accept_multiple_files=True)
    if archivos_pdf:
        st.write("### 1. Ordenar Archivos Subidos")
        nombres_archivos = [f.name for f in archivos_pdf]
        orden_archivos = st.multiselect("Orden de los documentos:", options=nombres_archivos, default=nombres_archivos)
        
        archivos_ordenados = []
        for name in orden_archivos:
            for f in archivos_pdf:
                if f.name == name and f not in archivos_ordenados:
                    archivos_ordenados.append(f)
                    break

        st.write("### 2. Ordenar páginas internas (Opcional)")
        orden_hojas = st.text_input("Reorganizar todas las hojas finales (Ej: 1,3,2,4,5). Deja en blanco para mantener el orden.")
        if st.button("Unir y Procesar"):
            with st.spinner("Uniendo documentos..."):
                doc_final = fitz.open()
                for pdf_file in archivos_ordenados:
                    doc_temp = fitz.open(stream=pdf_file.getvalue(), filetype="pdf")
                    doc_final.insert_pdf(doc_temp)
                    doc_temp.close()
                if orden_hojas:
                    try:
                        nuevo_orden = [int(x.strip()) - 1 for x in orden_hojas.split(",")]
                        doc_final.select(nuevo_orden)
                    except:
                        st.error("Error en el formato de orden.")
                        st.stop()
                salida_pdf = BytesIO()
                doc_final.save(salida_pdf)
                doc_final.close()
                st.success("¡Documento procesado!")
                st.download_button(label="Descargar PDF Unido", data=salida_pdf.getvalue(), file_name="documento_unido.pdf", mime="application/pdf")

# ==========================================
# MÓDULO 3: DIVIDIR O QUITAR HOJAS
# ==========================================
elif opcion == "Dividir o Quitar Hojas":
    st.header("✂️ Dividir o Quitar Hojas de un PDF")
    archivo_pdf = st.file_uploader("Sube tu archivo PDF", type=["pdf"])
    if archivo_pdf:
        doc = fitz.open(stream=archivo_pdf.getvalue(), filetype="pdf")
        num_paginas = len(doc)
        st.info(f"El documento tiene **{num_paginas} páginas** (numeradas del 1 al {num_paginas}).")
        
        modo = st.radio("¿Qué deseas hacer?", ["Conservar páginas (Extraer)", "Eliminar páginas (Quitar)"])
        paginas_input = st.text_input("Ingresa las páginas separadas por comas o rangos (Ej: 1, 3, 5-8):")
        
        if st.button("Procesar Documento"):
            with st.spinner("Procesando páginas..."):
                paginas_seleccionadas = set()
                for part in paginas_input.split(','):
                    part = part.strip()
                    if '-' in part:
                        try:
                            start, end = map(int, part.split('-'))
                            paginas_seleccionadas.update(range(start, end + 1))
                        except: pass
                    else:
                        try:
                            if part: paginas_seleccionadas.add(int(part))
                        except: pass
                
                idx_seleccionados = sorted([p - 1 for p in paginas_seleccionadas if 1 <= p <= num_paginas])
                idx_finales = [i for i in range(num_paginas) if i not in idx_seleccionados] if modo == "Eliminar páginas (Quitar)" else idx_seleccionados
                
                if not idx_finales:
                    st.error("La selección dejó el documento sin páginas.")
                else:
                    doc.select(idx_finales)
                    salida_pdf = BytesIO()
                    doc.save(salida_pdf)
                    doc.close()
                    st.success("¡Páginas procesadas con éxito!")
                    st.download_button(label="Descargar PDF Modificado", data=salida_pdf.getvalue(), file_name=f"modificado_{archivo_pdf.name}", mime="application/pdf")

# ==========================================
# MÓDULO 4: FIRMAR PDF O WORD (VISUAL Y DIGITAL)
# ==========================================
elif opcion == "Firmar PDF":
    import datetime
    import uuid
    import subprocess
    import json
    import base64
    from io import BytesIO
    from PIL import Image
    import traceback
    
    st.header("✍️ Firmar Documento (Soporta PDF y Word)")
    
    # --- INICIALIZACIÓN DE VARIABLES PARA PLANTILLAS (Evita errores de sesión) ---
    def init_keys(prefix):
        defaults = {
            f"check_texto_{prefix}": False if prefix == "vis" else True,
            f"nom_{prefix}": "David Moisés Lagos Riaño",
            f"mot_{prefix}": "Soy el autor de este documento" if prefix=="dig" else "Aprobación de documento",
            f"ciu_{prefix}": "Bogotá, Colombia",
            f"cor_{prefix}": "davidmlrdvm@gmail.com",
            f"tel_{prefix}": "3007351747",
            f"fec_{prefix}": True,
            f"font_{prefix}": 3.0,
            f"ancho_{prefix}": 280.0,
            f"alto_{prefix}": 80.0,
            f"fondo_{prefix}": True,
            f"img_perfil_{prefix}": None
        }
        for k, v in defaults.items():
            if k not in st.session_state:
                st.session_state[k] = v

    init_keys("vis")
    init_keys("dig")

    # Inicializar variables de estado para sincronizar selectores de página
    if "pag_vis" not in st.session_state: st.session_state.pag_vis = 1
    def sync_pag_vis_top(): st.session_state.pag_vis = st.session_state.pag_vis_top
    def sync_pag_vis_bot(): st.session_state.pag_vis = st.session_state.pag_vis_bot

    if "pag_dig" not in st.session_state: st.session_state.pag_dig = 1
    def sync_pag_dig_top(): st.session_state.pag_dig = st.session_state.pag_dig_top
    def sync_pag_dig_bot(): st.session_state.pag_dig = st.session_state.pag_dig_bot
    
    # --- MOTOR DE CONVERSIÓN Y CACHÉ ---
    def procesar_archivo_subido(archivo, key_prefix):
        session_key = f"cache_pdf_{key_prefix}_{archivo.name}"
        nombre_pdf = archivo.name.lower().replace(".docx", ".pdf")
        
        if session_key in st.session_state:
            return st.session_state[session_key], nombre_pdf
            
        if archivo.name.lower().endswith(".docx"):
            with st.spinner("Convirtiendo Word a PDF en el servidor..."):
                try:
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp_docx:
                        tmp_docx.write(archivo.getvalue())
                        ruta_docx = tmp_docx.name
                    
                    outdir = os.path.dirname(ruta_docx)
                    subprocess.run(
                        ["libreoffice", "--headless", "--convert-to", "pdf", ruta_docx, "--outdir", outdir],
                        check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE
                    )
                    ruta_pdf = ruta_docx.replace(".docx", ".pdf")
                    
                    with open(ruta_pdf, "rb") as f:
                        pdf_result = f.read()
                    
                    os.remove(ruta_docx)
                    os.remove(ruta_pdf)
                    
                    st.session_state[session_key] = pdf_result
                    return pdf_result, nombre_pdf
                except Exception as e:
                    st.error("Falló la conversión del documento Word.")
                    st.info("💡 IMPORTANTE: Si estás en Streamlit Cloud, asegúrate de tener el archivo `packages.txt` con la palabra `libreoffice`.")
                    st.stop()
        else:
            pdf_result = archivo.getvalue()
            st.session_state[session_key] = pdf_result
            return pdf_result, nombre_pdf

    # --- FUNCIÓN MAESTRA DE DIBUJO ---
    def dibujar_bloque_firma(page, rect, img_bytes, datos, font_size):
        if datos and any(datos.values()): 
            if img_bytes:
                img_rect = fitz.Rect(rect.x0 + 5, rect.y0 + 5, rect.x0 + (rect.width * 0.35), rect.y1 - 5)
                text_x0 = rect.x0 + (rect.width * 0.38)
                page.insert_image(img_rect, stream=img_bytes, keep_proportion=True)
            else:
                text_x0 = rect.x0 + 10
            
            lineas = []
            if datos.get("nombre"): lineas.append(f"Firmado por: {datos['nombre']}")
            if datos.get("motivo"): lineas.append(f"Motivo: {datos['motivo']}")
            if datos.get("ciudad"): lineas.append(f"Ubicación: {datos['ciudad']}")
            
            contacto = []
            if datos.get("correo"): contacto.append(datos['correo'])
            if datos.get("telefono"): contacto.append(datos['telefono'])
            if contacto: lineas.append(f"Contacto: {' | '.join(contacto)}")
            
            if datos.get("fecha"): 
                lineas.append(f"Fecha: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            
            y_text = rect.y0 + font_size + 4 
            espaciado = font_size * 1.4
            
            for linea in lineas:
                page.insert_text(fitz.Point(text_x0, y_text), linea, fontsize=font_size, color=(0, 0, 0), fontname="helv")
                y_text += espaciado
        else:
            if img_bytes:
                page.insert_image(rect, stream=img_bytes, keep_proportion=False)

    tipo_firma = st.radio("Elige el tipo de firma:", ["Firma Visual (Imagen y/o Texto en el documento)", "Firma Digital Criptográfica (Archivo .p12 / .pfx)"])
    
    # ---------------------------------------------------------
    # RAMA 1: FIRMA VISUAL
    # ---------------------------------------------------------
    if tipo_firma == "Firma Visual (Imagen y/o Texto en el documento)":
        archivo_subido = st.file_uploader("Paso 1: Sube el documento a firmar (PDF o Word)", type=["pdf", "docx"], key="pdf_firma_vis")
        if archivo_subido:
            pdf_bytes, nombre_pdf = procesar_archivo_subido(archivo_subido, "vis")
            doc_info = fitz.open(stream=pdf_bytes, filetype="pdf")
            num_paginas = len(doc_info)
            
            if st.session_state.pag_vis > num_paginas: st.session_state.pag_vis = num_paginas
            pagina_firma = st.session_state.pag_vis
            
            page_info = doc_info[pagina_firma - 1]
            max_width = float(page_info.rect.width)
            max_height = float(page_info.rect.height)
            doc_info.close()

            st.write("---")
            st.checkbox("📝 Añadir datos de texto junto a la rúbrica", key="check_texto_vis")
            
            datos_texto_firma = {}
            if st.session_state.check_texto_vis:
                st.write("### Datos del Firmante")
                col_d1, col_d2 = st.columns(2)
                with col_d1:
                    st.text_input("Nombre completo:", key="nom_vis")
                    st.text_input("Motivo de la firma:", key="mot_vis")
                    st.text_input("Ciudad / Ubicación:", key="ciu_vis")
                with col_d2:
                    st.text_input("Correo electrónico:", key="cor_vis")
                    st.text_input("Número de teléfono:", key="tel_vis")
                    st.checkbox("Incluir Fecha y Hora actual", key="fec_vis")
                    st.slider("Tamaño de la letra (Fuente)", 2.0, 16.0, step=0.5, key="font_vis")

                datos_texto_firma = {
                    "nombre": st.session_state.nom_vis,
                    "motivo": st.session_state.mot_vis,
                    "ciudad": st.session_state.ciu_vis,
                    "correo": st.session_state.cor_vis,
                    "telefono": st.session_state.tel_vis,
                    "fecha": st.session_state.fec_vis
                }

            archivo_firma = st.file_uploader("Sube la imagen de tu rúbrica (Opcional si usas texto)", type=["png", "jpg", "jpeg"], key="img_firma_vis")
            firma_bytes = None
            
            img_source = archivo_firma if archivo_firma else (BytesIO(st.session_state.img_perfil_vis) if st.session_state.img_perfil_vis else None)
            
            if img_source:
                firma_img = Image.open(img_source).convert("RGBA")
                st.checkbox("Quitar fondo blanco de la firma", key="fondo_vis")
                
                if st.session_state.fondo_vis:
                    datas = firma_img.getdata()
                    newData = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) or (item[0]>200 and item[1]>200 and item[2]>200) else item for item in datas]
                    firma_img.putdata(newData)
                
                img_byte_arr = BytesIO()
                firma_img.save(img_byte_arr, format='PNG')
                firma_bytes = img_byte_arr.getvalue()
                
            st.write("### Ajustes del Bloque / Imagen")
            col_t1, col_t2 = st.columns(2)
            with col_t1: ancho_firma = st.slider("Ancho", 10.0, 600.0, key="ancho_vis")
            with col_t2: alto_firma = st.slider("Alto", 10.0, 300.0, key="alto_vis")
                
            col_p1, col_p2 = st.columns(2)
            with col_p1: st.slider("Horizontal (X)", 0.0, max_width, value=st.session_state.firma_x, key="slider_firma_x", on_change=update_firma_pos)
            with col_p2: st.slider("Vertical (Y)", 0.0, max_height, value=st.session_state.firma_y, key="slider_firma_y", on_change=update_firma_pos)

            rect_firma = fitz.Rect(st.session_state.firma_x, st.session_state.firma_y, st.session_state.firma_x + ancho_firma, st.session_state.firma_y + alto_firma)
            
            st.write("### Vista Previa Interactiva")
            
            st.number_input("Página a firmar:", min_value=1, max_value=max(1, num_paginas), value=pagina_firma, key="pag_vis_top", on_change=sync_pag_vis_top)
            
            doc_preview = fitz.open(stream=pdf_bytes, filetype="pdf")
            page_preview = doc_preview[pagina_firma - 1]
            
            dibujar_bloque_firma(page_preview, rect_firma, firma_bytes, datos_texto_firma, st.session_state.font_vis)
            
            pix = page_preview.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
            img_pil = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            
            coords = streamlit_image_coordinates(img_pil, key="click_firma_vis")
            if coords is not None:
                click_id = f"{coords['x']}-{coords['y']}"
                if st.session_state.last_click_firma != click_id:
                    st.session_state.last_click_firma = click_id
                    st.session_state.firma_x = float(coords["x"]) / 1.5
                    st.session_state.firma_y = float(coords["y"]) / 1.5
                    st.rerun()
            doc_preview.close()
            
            st.number_input("Página a firmar:", min_value=1, max_value=max(1, num_paginas), value=pagina_firma, key="pag_vis_bot", on_change=sync_pag_vis_bot)
            
            # --- PANEL DE GUARDAR / CARGAR PLANTILLA (VISUAL) ---
            st.write("---")
            st.write("### ⚙️ Plantilla de Firma (Guardar / Cargar)")
            col_load, col_save = st.columns(2)
            with col_load:
                with st.expander("📁 Cargar plantilla guardada (.json)"):
                    perfil_subido = st.file_uploader("Sube tu archivo .json", type=["json"], key="load_prof_vis")
                    if st.button("Aplicar Plantilla", key="btn_apply_vis") and perfil_subido:
                        try:
                            texto_json = perfil_subido.getvalue().decode("utf-8")
                            data = json.loads(texto_json)
                            
                            st.session_state.check_texto_vis = bool(data.get("incluir_texto", st.session_state.check_texto_vis))
                            st.session_state.nom_vis = str(data.get("nombre", st.session_state.nom_vis))
                            st.session_state.mot_vis = str(data.get("motivo", st.session_state.mot_vis))
                            st.session_state.ciu_vis = str(data.get("ciudad", st.session_state.ciu_vis))
                            st.session_state.cor_vis = str(data.get("correo", st.session_state.cor_vis))
                            st.session_state.tel_vis = str(data.get("telefono", st.session_state.tel_vis))
                            st.session_state.fec_vis = bool(data.get("fecha", st.session_state.fec_vis))
                            st.session_state.font_vis = float(data.get("font_size", st.session_state.font_vis))
                            st.session_state.ancho_vis = float(data.get("ancho", st.session_state.ancho_vis))
                            st.session_state.alto_vis = float(data.get("alto", st.session_state.alto_vis))
                            st.session_state.fondo_vis = bool(data.get("quitar_fondo", st.session_state.fondo_vis))
                            
                            if data.get("imagen_b64"):
                                st.session_state.img_perfil_vis = base64.b64decode(data["imagen_b64"])
                            else:
                                st.session_state.img_perfil_vis = None
                                
                            st.success("Plantilla aplicada con éxito. Actualizando vista...")
                            st.rerun()
                        except Exception as e:
                            st.error(f"Error técnico al cargar: {str(e)}")
            with col_save:
                with st.expander("💾 Guardar configuración actual"):
                    img_to_save = archivo_firma.getvalue() if archivo_firma else st.session_state.img_perfil_vis
                    perfil_data = {
                        "incluir_texto": st.session_state.check_texto_vis,
                        "nombre": st.session_state.nom_vis,
                        "motivo": st.session_state.mot_vis,
                        "ciudad": st.session_state.ciu_vis,
                        "correo": st.session_state.cor_vis,
                        "telefono": st.session_state.tel_vis,
                        "fecha": st.session_state.fec_vis,
                        "font_size": st.session_state.font_vis,
                        "ancho": st.session_state.ancho_vis,
                        "alto": st.session_state.alto_vis,
                        "quitar_fondo": st.session_state.fondo_vis,
                        "imagen_b64": base64.b64encode(img_to_save).decode("utf-8") if img_to_save else None
                    }
                    st.download_button(
                        label="⬇️ Descargar plantilla (.json)",
                        data=json.dumps(perfil_data, indent=4),
                        file_name="plantilla_firma_visual.json",
                        mime="application/json",
                        key="down_prof_vis"
                    )

            st.write("---")
            st.write("### Opciones de Seguridad")
            bloquear_pdf = st.checkbox("🔒 Bloquear PDF contra modificaciones (Impide que se borre la firma)", key="block_vis")

            if st.button("Aplicar Firma y Descargar"):
                if not firma_bytes and not st.session_state.check_texto_vis:
                    st.warning("Debes subir una imagen o habilitar el texto para poder firmar.")
                else:
                    with st.spinner("Procesando documento final..."):
                        doc_final = fitz.open(stream=pdf_bytes, filetype="pdf")
                        page_final = doc_final[pagina_firma - 1]
                        
                        dibujar_bloque_firma(page_final, rect_firma, firma_bytes, datos_texto_firma, st.session_state.font_vis)
                        
                        salida_pdf = BytesIO()
                        if bloquear_pdf:
                            owner_pw = ''.join(random.choices(string.ascii_letters + string.digits, k=16))
                            perms = fitz.PDF_PERM_PRINT | fitz.PDF_PERM_ACCESSIBILITY
                            doc_final.save(salida_pdf, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=owner_pw, permissions=perms)
                        else:
                            doc_final.save(salida_pdf)
                        doc_final.close()
                        st.success("¡Firma aplicada con éxito!")
                        st.download_button(label="Descargar Documento Firmado", data=salida_pdf.getvalue(), file_name=f"firmado_{nombre_pdf}", mime="application/pdf")
    
    # ---------------------------------------------------------
    # RAMA 2: FIRMA DIGITAL CRIPTOGRÁFICA
    # ---------------------------------------------------------
    else:
        st.info("Aplica una firma digital con validez legal subiendo tu certificado PKCS#12 (.p12 o .pfx).")
        
        col_firma1, col_firma2 = st.columns(2)
        with col_firma1:
            archivo_subido = st.file_uploader("1. Sube el documento a firmar (PDF o Word)", type=["pdf", "docx"], key="pdf_crypto")
        with col_firma2:
            archivo_p12 = st.file_uploader("2. Sube tu certificado (.p12 / .pfx)", type=["p12", "pfx"])
        
        pdf_bytes = None
        nombre_pdf = ""
        if archivo_subido:
            pdf_bytes, nombre_pdf = procesar_archivo_subido(archivo_subido, "dig")

        password_certificado = st.text_input("3. Contraseña del certificado:", type="password")

        st.checkbox("✒️ Añadir bloque de firma profesional (Rúbrica + Datos textuales)", key="check_texto_dig")
        firma_bytes = None
        rect_firma = None
        pagina_firma = 1
        datos_texto_firma = {}

        if st.session_state.check_texto_dig and pdf_bytes:
            
            doc_info = fitz.open(stream=pdf_bytes, filetype="pdf")
            num_paginas = len(doc_info)
            if st.session_state.pag_dig > num_paginas: st.session_state.pag_dig = num_paginas
            pagina_firma = st.session_state.pag_dig
            
            page_info = doc_info[pagina_firma - 1]
            max_width = float(page_info.rect.width)
            max_height = float(page_info.rect.height)
            doc_info.close()
            
            st.write("---")
            st.write("### Datos del Firmante")
            col_d1, col_d2 = st.columns(2)
            with col_d1:
                st.text_input("Nombre completo:", key="nom_dig")
                st.text_input("Motivo de la firma:", key="mot_dig")
                st.text_input("Ciudad / Ubicación:", key="ciu_dig")
            with col_d2:
                st.text_input("Correo electrónico:", key="cor_dig")
                st.text_input("Número de teléfono:", key="tel_dig")
                st.checkbox("Incluir Fecha y Hora actual", key="fec_dig")
                st.slider("Tamaño de la letra (Fuente)", 2.0, 16.0, step=0.5, key="font_dig")

            datos_texto_firma = {
                "nombre": st.session_state.nom_dig,
                "motivo": st.session_state.mot_dig,
                "ciudad": st.session_state.ciu_dig,
                "correo": st.session_state.cor_dig,
                "telefono": st.session_state.tel_dig,
                "fecha": st.session_state.fec_dig
            }

            archivo_firma = st.file_uploader("Sube la imagen de tu rúbrica (Recomendado)", type=["png", "jpg", "jpeg"], key="img_firma_dig")
            
            img_source = archivo_firma if archivo_firma else (BytesIO(st.session_state.img_perfil_dig) if st.session_state.img_perfil_dig else None)
            
            if img_source:
                firma_img = Image.open(img_source).convert("RGBA")
                st.checkbox("Quitar fondo blanco", key="fondo_dig")
                
                if st.session_state.fondo_dig:
                    datas = firma_img.getdata()
                    newData = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) or (item[0]>200 and item[1]>200 and item[2]>200) else item for item in datas]
                    firma_img.putdata(newData)
                
                img_byte_arr = BytesIO()
                firma_img.save(img_byte_arr, format='PNG')
                firma_bytes = img_byte_arr.getvalue()
            
            st.write("### Ajuste del Bloque de Firma")
            col_t1, col_t2 = st.columns(2)
            with col_t1: ancho_firma = st.slider("Ancho del bloque", 10.0, 600.0, key="ancho_dig")
            with col_t2: alto_firma = st.slider("Alto del bloque", 10.0, 300.0, key="alto_dig")
                
            col_p1, col_p2 = st.columns(2)
            with col_p1: st.slider("Horizontal (X)", 0.0, max_width, value=st.session_state.firma_x, key="slider_firma_x", on_change=update_firma_pos)
            with col_p2: st.slider("Vertical (Y)", 0.0, max_height, value=st.session_state.firma_y, key="slider_firma_y", on_change=update_firma_pos)

            rect_firma = fitz.Rect(st.session_state.firma_x, st.session_state.firma_y, st.session_state.firma_x + ancho_firma, st.session_state.firma_y + alto_firma)

            st.write("### Vista Previa Interactiva")
            
            st.number_input("Página a firmar:", min_value=1, max_value=max(1, num_paginas), value=pagina_firma, key="pag_dig_top", on_change=sync_pag_dig_top)
            
            doc_preview = fitz.open(stream=pdf_bytes, filetype="pdf")
            page_preview = doc_preview[pagina_firma - 1]
            
            dibujar_bloque_firma(page_preview, rect_firma, firma_bytes, datos_texto_firma, st.session_state.font_dig)
            
            pix = page_preview.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
            img_pil = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            
            coords = streamlit_image_coordinates(img_pil, key="click_crypto_prev")
            if coords is not None:
                click_id = f"{coords['x']}-{coords['y']}"
                if st.session_state.last_click_firma != click_id:
                    st.session_state.last_click_firma = click_id
                    st.session_state.firma_x = float(coords["x"]) / 1.5
                    st.session_state.firma_y = float(coords["y"]) / 1.5
                    st.rerun()
            doc_preview.close()
            
            st.number_input("Página a firmar:", min_value=1, max_value=max(1, num_paginas), value=pagina_firma, key="pag_dig_bot", on_change=sync_pag_dig_bot)
            
            # --- PANEL DE GUARDAR / CARGAR PLANTILLA (DIGITAL) ---
            st.write("---")
            st.write("### ⚙️ Plantilla de Firma (Guardar / Cargar)")
            col_load, col_save = st.columns(2)
            with col_load:
                with st.expander("📁 Cargar plantilla guardada (.json)"):
                    perfil_subido_dig = st.file_uploader("Sube tu archivo .json", type=["json"], key="load_prof_dig")
                    if st.button("Aplicar Plantilla", key="btn_apply_dig") and perfil_subido_dig:
                        try:
                            texto_json = perfil_subido_dig.getvalue().decode("utf-8")
                            data = json.loads(texto_json)
                            
                            st.session_state.check_texto_dig = bool(data.get("incluir_texto", st.session_state.check_texto_dig))
                            st.session_state.nom_dig = str(data.get("nombre", st.session_state.nom_dig))
                            st.session_state.mot_dig = str(data.get("motivo", st.session_state.mot_dig))
                            st.session_state.ciu_dig = str(data.get("ciudad", st.session_state.ciu_dig))
                            st.session_state.cor_dig = str(data.get("correo", st.session_state.cor_dig))
                            st.session_state.tel_dig = str(data.get("telefono", st.session_state.tel_dig))
                            st.session_state.fec_dig = bool(data.get("fecha", st.session_state.fec_dig))
                            st.session_state.font_dig = float(data.get("font_size", st.session_state.font_dig))
                            st.session_state.ancho_dig = float(data.get("ancho", st.session_state.ancho_dig))
                            st.session_state.alto_dig = float(data.get("alto", st.session_state.alto_dig))
                            st.session_state.fondo_dig = bool(data.get("quitar_fondo", st.session_state.fondo_dig))
                            
                            if data.get("imagen_b64"):
                                st.session_state.img_perfil_dig = base64.b64decode(data["imagen_b64"])
                            else:
                                st.session_state.img_perfil_dig = None
                                
                            st.success("Plantilla aplicada con éxito. Actualizando vista...")
                            st.rerun()
                        except Exception as e:
                            st.error(f"Error técnico al cargar: {str(e)}")
            with col_save:
                with st.expander("💾 Guardar configuración actual"):
                    img_to_save = archivo_firma.getvalue() if archivo_firma else st.session_state.img_perfil_dig
                    perfil_data = {
                        "incluir_texto": st.session_state.check_texto_dig,
                        "nombre": st.session_state.nom_dig,
                        "motivo": st.session_state.mot_dig,
                        "ciudad": st.session_state.ciu_dig,
                        "correo": st.session_state.cor_dig,
                        "telefono": st.session_state.tel_dig,
                        "fecha": st.session_state.fec_dig,
                        "font_size": st.session_state.font_dig,
                        "ancho": st.session_state.ancho_dig,
                        "alto": st.session_state.alto_dig,
                        "quitar_fondo": st.session_state.fondo_dig,
                        "imagen_b64": base64.b64encode(img_to_save).decode("utf-8") if img_to_save else None
                    }
                    st.download_button(
                        label="⬇️ Descargar plantilla (.json)",
                        data=json.dumps(perfil_data, indent=4),
                        file_name="plantilla_firma_digital.json",
                        mime="application/json",
                        key="down_prof_dig"
                    )

            st.write("---")

        if archivo_p12 and password_certificado:
            if st.button("Diagnosticar Certificado"):
                try:
                    from cryptography.hazmat.primitives.serialization import pkcs12
                    private_key, cert, extras = pkcs12.load_key_and_certificates(
                        archivo_p12.getvalue(),
                        password_certificado.encode("utf-8")
                    )
                    st.success("Certificado cargado correctamente")
                    st.write("Clave privada:", private_key is not None)
                    st.write("Certificado:", cert is not None)
                    if cert:
                        st.code(str(cert.subject))
                except Exception as e:
                    st.error(f"ERROR REAL: {str(e)}")
                    st.code(traceback.format_exc())

        if archivo_subido and archivo_p12 and password_certificado:
            if st.button("Firmar Documento Digitalmente", type="primary"):
                with st.spinner("Procesando firma criptográfica..."):
                    ruta_p12 = None
                    try:
                        password_bytes = password_certificado.encode("utf-8")
                        
                        with tempfile.NamedTemporaryFile(delete=False, suffix=".p12") as tmp_p12:
                            tmp_p12.write(archivo_p12.getvalue())
                            ruta_p12 = tmp_p12.name
                            
                        signer = signers.SimpleSigner.load_pkcs12(
                            pfx_file=ruta_p12,
                            passphrase=password_bytes
                        )
                        
                        if st.session_state.check_texto_dig and rect_firma:
                            doc_visual = fitz.open(stream=pdf_bytes, filetype="pdf")
                            page_visual = doc_visual[pagina_firma - 1]
                            
                            dibujar_bloque_firma(page_visual, rect_firma, firma_bytes, datos_texto_firma, st.session_state.font_dig)
                            
                            pdf_base = BytesIO()
                            doc_visual.save(pdf_base)
                            doc_visual.close()
                        else:
                            pdf_base = BytesIO(pdf_bytes)
                        
                        pdf_writer = IncrementalPdfFileWriter(pdf_base)
                        campo_unico = f"FirmaDigitalLegal_{uuid.uuid4().hex[:8]}"
                        
                        signers.sign_pdf(
                            pdf_writer, 
                            signers.PdfSignatureMetadata(
                                field_name=campo_unico,
                                location=datos_texto_firma.get("ciudad", ""),
                                reason=datos_texto_firma.get("motivo", "")
                            ), 
                            signer=signer, 
                            in_place=True
                        )
                        
                        st.success("¡PDF firmado correctamente con bloque visual y certificado digital!")
                        st.download_button(
                            label="⬇️ Descargar Documento Firmado",
                            data=pdf_base.getvalue(),
                            file_name=f"firmado_digital_{nombre_pdf}",
                            mime="application/pdf"
                        )
                        
                    except Exception as e:
                        st.error(f"ERROR REAL: {str(e)}")
                        st.code(traceback.format_exc())
                    finally:
                        if ruta_p12 and os.path.exists(ruta_p12):
                            try:
                                os.remove(ruta_p12)
                            except Exception:
                                pass
                             
# ==========================================
# MÓDULO 5: EDITAR Y RELLENAR PDF
# ==========================================
elif opcion == "Editar y Rellenar PDF":
    st.header("📝 Rellenar Formularios PDF (Multicampo)")
    archivo_pdf = st.file_uploader("Sube el formulario PDF a rellenar", type=["pdf"])
    
    if archivo_pdf:
        if st.session_state.pdf_actual_form != archivo_pdf.name:
            st.session_state.textos_agregados = []
            st.session_state.pdf_actual_form = archivo_pdf.name
            st.session_state.texto_x = 50.0
            st.session_state.texto_y = 50.0
            st.session_state.last_click_texto = None

        pdf_bytes = archivo_pdf.getvalue()
        doc_info = fitz.open(stream=pdf_bytes, filetype="pdf")
        num_paginas = len(doc_info)
        
        col_herramientas, col_vista = st.columns([1, 2], gap="large")

        with col_herramientas:
            st.write("### 1. Herramientas de Texto")
            pagina_texto = st.number_input("Página actual:", min_value=1, max_value=num_paginas, value=1)
            texto_insertar = st.text_input("Escribe aquí el texto a insertar:", "Ej: Carlos Gómez")
            
            col_f1, col_f2 = st.columns(2)
            with col_f1: tamano_fuente = st.slider("Tamaño letra", 8, 72, 12)
            with col_f2: color_texto = st.color_picker("Color", "#000000")
            
            rgb_color = tuple(int(color_texto.lstrip('#')[i:i+2], 16)/255 for i in (0, 2, 4))
            page_info = doc_info[pagina_texto - 1]
            max_width = float(page_info.rect.width)
            max_height = float(page_info.rect.height)
            
            st.slider("Ajuste Fino X", 0.0, max_width, value=st.session_state.texto_x, key="slider_texto_x", on_change=update_texto_pos)
            st.slider("Ajuste Fino Y", 0.0, max_height, value=st.session_state.texto_y, key="slider_texto_y", on_change=update_texto_pos)
            
            if st.button("➕ Añadir este texto al documento", type="primary", use_container_width=True):
                if texto_insertar.strip():
                    st.session_state.textos_agregados.append({
                        "pagina": pagina_texto, "texto": texto_insertar,
                        "x": st.session_state.texto_x, "y": st.session_state.texto_y,
                        "size": tamano_fuente, "color": rgb_color
                    })
                    st.session_state.texto_y = min(max_height, st.session_state.texto_y + 25.0)
                    st.rerun()

            st.write("---")
            st.write("### 📋 Textos Agregados")
            if st.session_state.textos_agregados:
                for i, txt in enumerate(st.session_state.textos_agregados):
                    c1, c2 = st.columns([4, 1])
                    c1.write(f"**Pg {txt['pagina']}:** {txt['texto']}")
                    if c2.button("🗑️", key=f"del_{i}"):
                        st.session_state.textos_agregados.pop(i)
                        st.rerun()
                
                st.write("---")
                if st.button("⬇️ Aplicar Todos los Textos y Descargar", type="primary"):
                    with st.spinner("Procesando..."):
                        doc_final = fitz.open(stream=pdf_bytes, filetype="pdf")
                        for txt in st.session_state.textos_agregados:
                            page_final = doc_final[txt["pagina"] - 1]
                            p_final = fitz.Point(txt["x"], txt["y"])
                            page_final.insert_text(p_final, txt["texto"], fontsize=txt["size"], color=txt["color"])
                        
                        salida_pdf = BytesIO()
                        doc_final.save(salida_pdf)
                        doc_final.close()
                        st.success("¡Formulario completado!")
                        st.download_button(label="Descargar PDF Final", data=salida_pdf.getvalue(), file_name=f"formulario_{archivo_pdf.name}", mime="application/pdf")
            else:
                st.info("Aún no has agregado ningún texto.")
        
        with col_vista:
            st.write("### Vista Previa Interactiva (Clic para posicionar)")
            doc_preview = fitz.open(stream=pdf_bytes, filetype="pdf")
            page_preview = doc_preview[pagina_texto - 1]
            
            for txt in st.session_state.textos_agregados:
                if txt["pagina"] == pagina_texto:
                    p_hist = fitz.Point(txt["x"], txt["y"])
                    page_preview.insert_text(p_hist, txt["texto"], fontsize=txt["size"], color=txt["color"])
            
            if texto_insertar.strip():
                p_actual = fitz.Point(st.session_state.texto_x, st.session_state.texto_y)
                page_preview.insert_text(p_actual, texto_insertar, fontsize=tamano_fuente, color=rgb_color)
            
            pix = page_preview.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
            img_pil = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            
            coords = streamlit_image_coordinates(img_pil, key="click_texto")
            if coords is not None:
                click_id = f"{coords['x']}-{coords['y']}"
                if st.session_state.last_click_texto != click_id:
                    st.session_state.last_click_texto = click_id
                    st.session_state.texto_x = float(coords["x"]) / 1.5
                    st.session_state.texto_y = float(coords["y"]) / 1.5
                    st.rerun()

        doc_info.close()

# ==========================================
# MÓDULO 6: MARCA DE AGUA
# ==========================================
elif opcion == "Marca de Agua":
    st.header("©️ Insertar Marca de Agua")
    archivo_pdf = st.file_uploader("Sube el PDF", type=["pdf"])
    
    if archivo_pdf:
        pdf_bytes = archivo_pdf.getvalue()
        
        st.write("### Configuración de la Marca")
        tipo = st.radio("Tipo de marca:", ["Texto", "Imagen (Logo)"])
        
        if tipo == "Texto":
            texto_marca = st.text_input("Texto de la marca de agua:", "CONFIDENCIAL")
            col1, col2, col3 = st.columns(3)
            with col1: color_marca = st.color_picker("Color del texto", "#808080")
            with col2: size_marca = st.slider("Tamaño de texto", 20, 150, 60)
            with col3: opacidad = st.slider("Opacidad", 0.1, 1.0, 0.5)
            
            rgb_color = tuple(int(color_marca.lstrip('#')[i:i+2], 16)/255 for i in (0, 2, 4))
            
            if st.button("Aplicar Marca de Agua (Todas las hojas)"):
                with st.spinner("Aplicando..."):
                    doc_final = fitz.open(stream=pdf_bytes, filetype="pdf")
                    for page in doc_final:
                        rect = page.rect
                        x_center = rect.width / 2 - (len(texto_marca) * size_marca / 4)
                        y_center = rect.height / 2
                        p = fitz.Point(x_center, y_center)
                        page.insert_text(p, texto_marca, fontsize=size_marca, color=rgb_color, fill_opacity=opacidad)
                    
                    salida_pdf = BytesIO()
                    doc_final.save(salida_pdf)
                    doc_final.close()
                    st.success("¡Marca de agua de texto aplicada!")
                    st.download_button("Descargar PDF con Marca", data=salida_pdf.getvalue(), file_name=f"marca_agua_{archivo_pdf.name}", mime="application/pdf")
        
        else:
            img_marca = st.file_uploader("Sube tu logo/imagen", type=["png", "jpg", "jpeg"])
            if img_marca:
                col1, col2 = st.columns(2)
                with col1: escala = st.slider("Tamaño de la imagen (%)", 10, 100, 50)
                with col2: opacidad_img = st.slider("Opacidad", 0.1, 1.0, 0.3)
                
                if st.button("Aplicar Logo como Marca de Agua"):
                    with st.spinner("Procesando..."):
                        img = Image.open(img_marca).convert("RGBA")
                        alpha = img.split()[3]
                        alpha = ImageEnhance.Brightness(alpha).enhance(opacidad_img)
                        img.putalpha(alpha)
                        
                        img_byte_arr = BytesIO()
                        img.save(img_byte_arr, format='PNG')
                        img_bytes_transparente = img_byte_arr.getvalue()
                        
                        doc_final = fitz.open(stream=pdf_bytes, filetype="pdf")
                        for page in doc_final:
                            rect = page.rect
                            img_w = (rect.width * escala) / 100
                            img_h = (img.height * img_w) / img.width
                            
                            x0 = (rect.width - img_w) / 2
                            y0 = (rect.height - img_h) / 2
                            rect_img = fitz.Rect(x0, y0, x0 + img_w, y0 + img_h)
                            
                            page.insert_image(rect_img, stream=img_bytes_transparente)
                            
                        salida_pdf = BytesIO()
                        doc_final.save(salida_pdf)
                        doc_final.close()
                        st.success("¡Logo aplicado como marca de agua!")
                        st.download_button("Descargar PDF con Logo", data=salida_pdf.getvalue(), file_name=f"logo_agua_{archivo_pdf.name}", mime="application/pdf")

# ==========================================
# MÓDULO 7: PROTEGER / DESBLOQUEAR
# ==========================================
elif opcion == "Proteger / Desbloquear PDF":
    st.header("🔐 Proteger o Desbloquear PDF")
    modo_seguridad = st.radio("Elige una acción:", ["Proteger (Añadir Contraseña)", "Desbloquear (Quitar Contraseña)"])
    
    archivo_pdf = st.file_uploader(f"Sube el PDF a {modo_seguridad.split(' ')[0].lower()}", type=["pdf"])
    
    if archivo_pdf:
        pdf_bytes = archivo_pdf.getvalue()
        
        if "Proteger" in modo_seguridad:
            col1, col2 = st.columns(2)
            with col1: user_pass = st.text_input("Contraseña de Apertura (Para leer):", type="password")
            with col2: owner_pass = st.text_input("Contraseña Maestra (Para editar/imprimir):", type="password")
            
            st.write("### Permisos para los Lectores")
            perm_imprimir = st.checkbox("Permitir Imprimir", value=True)
            perm_copiar = st.checkbox("Permitir Copiar Texto y Gráficos", value=True)
            perm_editar = st.checkbox("Permitir Modificar el Documento", value=False)
            
            if st.button("Proteger y Descargar"):
                if not user_pass and not owner_pass:
                    st.error("Debes ingresar al menos una contraseña.")
                else:
                    with st.spinner("Encriptando documento con AES-256..."):
                        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
                        perms = fitz.PDF_PERM_ACCESSIBILITY 
                        if perm_imprimir: perms |= fitz.PDF_PERM_PRINT
                        if perm_copiar: perms |= fitz.PDF_PERM_COPY
                        if perm_editar: perms |= fitz.PDF_PERM_MODIFY
                        
                        salida_pdf = BytesIO()
                        doc.save(salida_pdf, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=user_pass, owner_pw=owner_pass if owner_pass else user_pass, permissions=perms)
                        doc.close()
                        st.success("¡Documento protegido exitosamente!")
                        st.download_button("Descargar PDF Seguro", data=salida_pdf.getvalue(), file_name=f"protegido_{archivo_pdf.name}", mime="application/pdf")
        
        else:
            pass_actual = st.text_input("Ingresa la contraseña actual del documento:", type="password")
            if st.button("Desbloquear y Descargar"):
                with st.spinner("Intentando desbloquear..."):
                    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
                    if doc.is_encrypted:
                        if doc.authenticate(pass_actual):
                            salida_pdf = BytesIO()
                            doc.save(salida_pdf, encryption=fitz.PDF_ENCRYPT_NONE)
                            doc.close()
                            st.success("¡Documento desbloqueado exitosamente!")
                            st.download_button("Descargar PDF Liberado", data=salida_pdf.getvalue(), file_name=f"desbloqueado_{archivo_pdf.name}", mime="application/pdf")
                        else:
                            st.error("❌ Contraseña incorrecta. El documento sigue bloqueado.")
                            doc.close()
                    else:
                        st.warning("Este documento no está protegido por contraseña.")
                        doc.close()

# ==========================================
# MÓDULO 8: CONVERTIR ARCHIVOS
# ==========================================
elif opcion == "Convertir a Word, Excel, PPT":
    st.header("🔄 Convertir PDF a Otros Formatos")
    archivo_pdf = st.file_uploader("Sube tu archivo PDF", type=["pdf"])
    if archivo_pdf is not None:
        formato = st.radio("Selecciona el formato:", ["Word (.docx) - Texto editable", "Excel (.xlsx) - Extraer Tablas", "PowerPoint (.pptx) - Presentación visual"])
        if st.button(f"Convertir a {formato.split(' ')[0]}"):
            if "Word" in formato:
                with st.spinner("Convirtiendo a Word..."):
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_pdf:
                        tmp_pdf.write(archivo_pdf.getvalue())
                        ruta_pdf = tmp_pdf.name
                    ruta_docx = ruta_pdf.replace(".pdf", ".docx")
                    try:
                        cv = Converter(ruta_pdf)
                        cv.convert(ruta_docx, start=0, end=None)
                        cv.close()
                        with open(ruta_docx, "rb") as docx_file:
                            docx_bytes = docx_file.read()
                        st.success("¡Conversión exitosa!")
                        st.download_button(label="Descargar Word", data=docx_bytes, file_name=f"{archivo_pdf.name.replace('.pdf', '')}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                    finally:
                        if os.path.exists(ruta_pdf): os.remove(ruta_pdf)
                        if os.path.exists(ruta_docx): os.remove(ruta_docx)
            elif "Excel" in formato:
                with st.spinner("Escaneando tablas..."):
                    try:
                        todas_las_tablas = []
                        with pdfplumber.open(BytesIO(archivo_pdf.getvalue())) as pdf:
                            for page in pdf.pages:
                                tables = page.extract_tables()
                                for table in tables:
                                    if table: 
                                        df = pd.DataFrame(table[1:], columns=table[0]) if len(table)>1 else pd.DataFrame(table)
                                        todas_las_tablas.append(df)
                        if not todas_las_tablas: st.warning("No se detectaron tablas.")
                        else:
                            salida_excel = BytesIO()
                            with pd.ExcelWriter(salida_excel, engine='openpyxl') as writer:
                                for i, df in enumerate(todas_las_tablas): df.to_excel(writer, sheet_name=f"Tabla_{i+1}", index=False)
                            st.success(f"¡{len(todas_las_tablas)} tablas extraídas!")
                            st.download_button(label="Descargar Excel", data=salida_excel.getvalue(), file_name=f"{archivo_pdf.name.replace('.pdf', '')}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    except Exception as e: st.error(f"Error al extraer: {e}")
            elif "PowerPoint" in formato:
                with st.spinner("Creando diapositivas..."):
                    try:
                        prs = Presentation()
                        blank_slide_layout = prs.slide_layouts[6] 
                        doc = fitz.open(stream=archivo_pdf.getvalue(), filetype="pdf")
                        for page_num in range(len(doc)):
                            page = doc[page_num]
                            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2)) 
                            slide = prs.slides.add_slide(blank_slide_layout)
                            slide.shapes.add_picture(BytesIO(pix.tobytes("png")), 0, 0, width=prs.slide_width, height=prs.slide_height)
                        doc.close()
                        salida_ppt = BytesIO()
                        prs.save(salida_ppt)
                        st.success("¡Conversión exitosa!")
                        st.download_button(label="Descargar PowerPoint", data=salida_ppt.getvalue(), file_name=f"{archivo_pdf.name.replace('.pdf', '')}.pptx", mime="application/vnd.openxmlformats-officedocument.presentationml.presentation")
                    except Exception as e: st.error(f"Error creando PowerPoint: {e}")
